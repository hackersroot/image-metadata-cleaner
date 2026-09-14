"""Regression tests. Extra test dependencies: python-docx openpyxl python-pptx pymupdf."""
import csv
import queue
import struct
import tempfile
import threading
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch
import xml.etree.ElementTree as ET

import pymupdf as fitz
import olefile
from PIL import Image, PngImagePlugin
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pypdf import PdfReader, PdfWriter
from pypdf.generic import NameObject, DecodedStreamObject, DictionaryObject
from image_metadata_cleaner import App, clean_file, scan_files, export_report

SECRET = 'PRIVATE_METADATA_12345'


def make_ole(path, main):
    """Synthetic CFB fixture, not a full Word/Excel/PowerPoint document."""
    free, end = 0xFFFFFFFF, 0xFFFFFFFE
    header = bytearray(512)
    header[:8] = bytes.fromhex('D0CF11E0A1B11AE1')
    struct.pack_into('<HHHH', header, 24, 0x3e, 3, 0xfffe, 9)
    struct.pack_into('<H', header, 32, 6)
    struct.pack_into('<IIIIIIIII', header, 40, 0, 1, 1, 0, 4096, end, 0, end, 0)
    struct.pack_into('<109I', header, 76, 0, *([free]*108))
    fat = [free]*128; fat[0] = 0xFFFFFFFD; fat[1] = end
    for start in (2, 10, 18):
        for i in range(start, start+7): fat[i] = i+1
        fat[start+7] = end
    def entry(name, kind, start, size, right=free, child=free):
        data = bytearray(128); encoded = (name+'\0').encode('utf-16le')
        data[:len(encoded)] = encoded
        struct.pack_into('<HBBIII', data, 64, len(encoded), kind, 1, free, right, child)
        struct.pack_into('<IQ', data, 116, start, size)
        return data
    directory = entry('Root Entry',5,end,0,child=1)
    directory += entry(main,2,2,4096,right=2)
    directory += entry('\x05SummaryInformation',2,10,4096,right=3)
    directory += entry('\x05DocumentSummaryInformation',2,18,4096)
    prop = bytearray(4096); prop[:2] = b'\xfe\xff'
    struct.pack_into('<I', prop, 24, 1)
    prop[28:44] = bytes.fromhex('E0859FF2F94F6810AB9108002B27B3D9')
    struct.pack_into('<I', prop, 44, 48)
    value = (SECRET+'\0').encode()
    struct.pack_into('<IIIIII', prop, 48, 24+len(value), 1, 4, 16, 30, len(value))
    prop[72:72+len(value)] = value
    body = bytearray(4096); body[100:110] = b'KEEP_BODY!'
    path.write_bytes(header + struct.pack('<128I',*fat) + directory + body + prop + prop)


class CleanerTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
    def tearDown(self): self.temp.cleanup()

    def test_images(self):
        for fmt, ext in [('JPEG','jpg'),('PNG','png'),('WEBP','webp'),('BMP','bmp'),('TIFF','tiff')]:
            with self.subTest(format=fmt):
                src = self.root/f'in.{ext}'; dst = self.root/f'out.{ext}'
                im = Image.new('RGB',(12,8),'red'); exif = Image.Exif()
                exif[315] = SECRET; exif[274] = 6
                opts = {'exif':exif} if fmt != 'BMP' else {}
                if fmt=='PNG':
                    info = PngImagePlugin.PngInfo(); info.add_text('Author',SECRET); opts['pnginfo'] = info
                im.save(src,format=fmt,**opts); clean_file(src,dst)
                with Image.open(dst) as check:
                    self.assertNotIn(315,check.getexif())
                    self.assertEqual(check.size,(8,12) if fmt!='BMP' else (12,8))
                self.assertNotIn(SECRET.encode(),dst.read_bytes())

    def test_pdf_metadata_render_content_and_form(self):
        src = self.root/'in.pdf'; dst = self.root/'out.pdf'
        pdf = fitz.open(); page = pdf.new_page(); page.insert_text((72,72),'Document content survives')
        widget = fitz.Widget(); widget.field_name = 'Name'; widget.field_type = fitz.PDF_WIDGET_TYPE_TEXT
        widget.field_value = 'Alice'; widget.rect = fitz.Rect(72,100,220,125); page.add_widget(widget)
        pdf.set_metadata({'author':SECRET,'title':SECRET}); pdf.save(src); pdf.close()
        writer = PdfWriter(clone_from=src)
        xmp = DecodedStreamObject(); xmp.set_data(SECRET.encode())
        xmp[NameObject('/Type')] = NameObject('/Metadata')
        writer.root_object[NameObject('/Metadata')] = writer._add_object(xmp)
        writer.pages[0][NameObject('/Metadata')] = writer._add_object(xmp)
        writer.write(src)
        clean_file(src,dst)
        self.assertNotIn(SECRET.encode(),dst.read_bytes())
        self.assertFalse(PdfReader(dst).metadata)
        self.assertEqual(PdfReader(dst).get_fields()['Name']['/V'],'Alice')
        with fitz.open(src) as before, fitz.open(dst) as after:
            self.assertEqual(before[0].get_pixmap().samples,after[0].get_pixmap().samples)
            self.assertEqual(before[0].get_text(),after[0].get_text())

    def test_pdf_encrypted_and_signed_preserve_original(self):
        for signed in (False,True):
            src = self.root/'guard.pdf'; writer = PdfWriter(); writer.add_blank_page(100,100)
            if signed: writer.root_object[NameObject('/TestSignature')] = DictionaryObject({NameObject('/FT'):NameObject('/Sig')})
            else: writer.encrypt('password')
            writer.write(src); before = src.read_bytes()
            with self.assertRaises(ValueError): clean_file(src,src)
            self.assertEqual(before,src.read_bytes())

    def test_office_properties_and_content(self):
        for ext in ('docx','xlsx','pptx'):
            with self.subTest(format=ext):
                src = self.root/f'in.{ext}'; dst = self.root/f'out.{ext}'
                if ext=='docx':
                    doc = Document(); doc.add_paragraph('Keep document text'); doc.core_properties.author = SECRET; doc.save(src)
                elif ext=='xlsx':
                    doc = Workbook(); doc.active['A1'] = '=1+2'; doc.properties.creator = SECRET; doc.save(src)
                else:
                    doc = Presentation(); doc.slides.add_slide(doc.slide_layouts[0]).shapes.title.text = 'Keep title'
                    doc.core_properties.author = SECRET; doc.save(src)
                with zipfile.ZipFile(src,'a') as package:
                    package.writestr('docProps/custom.xml', '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties"><property name="Secret">'+SECRET+'</property></Properties>')
                clean_file(src,dst)
                with zipfile.ZipFile(src) as before, zipfile.ZipFile(dst) as after:
                    for name in before.namelist():
                        if not name.startswith('docProps/') and name not in ('_rels/.rels','[Content_Types].xml'):
                            self.assertEqual(before.read(name),after.read(name),name)
                    for name in after.namelist():
                        self.assertNotIn(SECRET.encode(),after.read(name))
                if ext=='docx': self.assertEqual(Document(dst).paragraphs[0].text,'Keep document text')
                elif ext=='xlsx': self.assertEqual(load_workbook(dst).active['A1'].value,'=1+2')
                else: self.assertEqual(Presentation(dst).slides[0].shapes.title.text,'Keep title')

    def test_legacy_property_streams(self):
        for ext, main in [('doc','WordDocument'),('xls','Workbook'),('ppt','PowerPoint Document')]:
            src = self.root/f'in.{ext}'; dst = self.root/f'out.{ext}'; make_ole(src,main)
            with olefile.OleFileIO(src) as before:
                self.assertEqual(before.getproperties('\x05SummaryInformation')[4].decode(),SECRET)
                body = before.openstream(main).read()
            clean_file(src,dst)
            with olefile.OleFileIO(dst) as after:
                self.assertEqual(after.getproperties('\x05SummaryInformation'),{})
                self.assertEqual(after.getproperties('\x05DocumentSummaryInformation'),{})
                self.assertEqual(body,after.openstream(main).read())
            self.assertNotIn(SECRET.encode(),dst.read_bytes())

    def test_transparency_animation_and_corruption(self):
        src = self.root/'pal.png'; dst = self.root/'out.png'
        im = Image.new('P',(4,4)); im.info['transparency'] = 0; im.save(src); clean_file(src,dst)
        with Image.open(dst) as check: self.assertEqual(check.convert('RGBA').getpixel((0,0))[3],0)
        frames = [Image.new('RGB',(4,4),c) for c in ['red','blue']]
        frames[0].save(src,save_all=True,append_images=frames[1:])
        with self.assertRaises(ValueError): clean_file(src,dst)
        src = self.root/'broken.pdf'; src.write_bytes(b'broken')
        with self.assertRaises(Exception): clean_file(src,src)
        self.assertEqual(src.read_bytes(),b'broken')

    def test_office_signed_guard(self):
        src = self.root/'signed.docx'; doc = Document(); doc.save(src)
        with zipfile.ZipFile(src,'a') as package: package.writestr('_xmlsignatures/sig1.xml','signature')
        before = src.read_bytes()
        with self.assertRaises(ValueError): clean_file(src,src)
        self.assertEqual(before,src.read_bytes())

    def test_gui_folder_zip_replace_and_report(self):
        source = self.root/'input'; source.mkdir(); (source/'nested').mkdir()
        Image.new('RGB',(8,8)).save(source/'nested'/'image.png')
        (source/'unknown.txt').write_text('unchanged')
        app = App(); app.withdraw(); app.update()
        try:
            for mode in ('Folder','ZIP','Replace'):
                target = self.root/'export.zip' if mode=='ZIP' else self.root
                app.run_batch(source,True,mode,target)
                app.poll()
                self.assertIn('1 cleaned',app.status.get(),app.status.get())
                self.assertIn('1 skipped',app.status.get())
            folders = list(self.root.glob('cleaned_files_*')); self.assertEqual(len(folders),1)
            self.assertTrue((folders[0]/'nested'/'image.png').exists())
            self.assertTrue((folders[0]/'metadata_report.csv').exists())
            with zipfile.ZipFile(self.root/'export.zip') as archive:
                self.assertIn('nested/image.png',archive.namelist())
                self.assertIn('metadata_report.csv',archive.namelist())
                self.assertIsNone(archive.testzip())
            report = self.root/'manual.csv'
            with patch('image_metadata_cleaner.filedialog.asksaveasfilename',return_value=str(report)):
                app.save_report()
            with report.open(encoding='utf-8-sig') as stream: self.assertEqual(len(list(csv.DictReader(stream))),2)
            app.cancel.set(); app.run_batch(source,True,'Folder',self.root); app.poll()
            self.assertIn('2 unprocessed',app.status.get())
        finally: app.destroy()

    def test_gui_start_worker(self):
        import time
        source = self.root/'source'; source.mkdir()
        Image.new('RGB',(8,8)).save(source/'image.png')
        app = App(); app.withdraw()
        try:
            app.folder.set(str(source)); app.destination.set(str(self.root))
            app.start()
            self.assertTrue(app.busy)
            deadline = time.monotonic()+10
            while app.busy and time.monotonic()<deadline:
                app.update(); time.sleep(0.01)
            self.assertFalse(app.busy)
            self.assertIn('1 cleaned',app.status.get())
            self.assertEqual(str(app.report['state']),'normal')
        finally: app.destroy()

    def test_csv_formula_escaping(self):
        report = self.root/'report.csv'
        export_report([{'file':'=1+1','status':'failed','detail':'@test'}],report)
        with report.open(encoding='utf-8-sig') as stream:
            row = next(csv.DictReader(stream)); self.assertEqual(row['file'],"'=1+1")

    def test_recursive_output_exclusion(self):
        (self.root/'cleaned_files_previous').mkdir()
        (self.root/'cleaned_files_previous'/'skip.png').touch()
        (self.root/'keep.png').touch()
        self.assertEqual([p.name for p in scan_files(self.root,True)],['keep.png'])


if __name__ == '__main__':
    unittest.main(verbosity=2)
